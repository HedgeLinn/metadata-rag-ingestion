"""实验编排器 — 多维度对比 + 完整/召回两种模式。"""
from __future__ import annotations

import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, asdict
from datetime import datetime
from pathlib import Path

import openpyxl

from config import DEFAULT_KB, DEFAULT_TOP_K, GEN_WORKERS, SCORE_WORKERS, SEARCH_WORKERS, PROMPTS
from core.registry import factory
from core.qdrant import get_client
from search import search as strategy_search
from pipeline.recall import calc_recall, calc_recall_merged, calc_recall_concat, calc_file_recall
from pipeline.generate import build_rag_chain


@dataclass
class ExperimentConfig:
    """一次实验的配置。variable 指定唯一可多选的对比维度。"""
    name: str = ""
    kb_name: str = DEFAULT_KB
    test_set_path: str = ""
    output_dir: str = ""

    variable: str = "chunker"           # chunker | retriever | reranker
    variable_values: list[str] = field(default_factory=list)

    chunker: str = "fixed_1100"
    retriever: str = "hybrid"
    reranker: str = "direct"
    top_k: int = DEFAULT_TOP_K
    score_model: str = "deepseek-v4-pro"
    score_mode: str = "rubric"
    runs: int = 1

    def __post_init__(self):
        if not self.name:
            self.name = f"完整实验_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        if not self.output_dir:
            self.output_dir = f"output/{self.kb_name}/{self.name}"


def expand_pipelines(config: ExperimentConfig) -> list[dict]:
    """沿对比维度展开，其余维度固定。"""
    pipelines = []
    for value in config.variable_values:
        pipe = {
            "label": value,
            "chunker": value if config.variable == "chunker" else config.chunker,
            "retriever": value if config.variable == "retriever" else config.retriever,
            "reranker": value if config.variable == "reranker" else config.reranker,
            "top_k": config.top_k,
        }
        pipelines.append(pipe)
    return pipelines


class ExperimentRunner:
    """执行完整实验：检索 → 生成 → 召回 → 评分 → 报告。"""

    def __init__(self, config: ExperimentConfig):
        self.config = config

    def run(self):
        config = self.config
        out_dir = Path(config.output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        (out_dir / "config.json").write_text(
            json.dumps(asdict(config), ensure_ascii=False, indent=2), encoding="utf-8")

        test_set = self._load_test_set(config.test_set_path)
        pipelines = expand_pipelines(config)

        print(f"实验: {config.name}")
        print(f"知识库: {config.kb_name}")
        print(f"对比维度: {config.variable}")
        print(f"组合数: {len(pipelines)}  |  测试题数: {len(test_set)}")
        print()

        self._ensure_built(pipelines)

        t_start = time.time()
        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "汇总"

        # 表头
        headers = ["问题ID", "问题", "正确回答", "point", "必答点", "选答点", "参考文件"]
        for ci, h in enumerate(headers, start=1):
            ws_summary.cell(row=1, column=ci, value=h)
        col = 8
        for pipe in pipelines:
            ws_summary.cell(row=1, column=col, value=pipe["label"])
            ws_summary.cell(row=1, column=col + 1, value=f"{pipe['label']}_召回")
            ws_summary.cell(row=1, column=col + 2, value=f"{pipe['label']}_召回位置")
            col += 3

        for qi, q in enumerate(test_set):
            row = qi + 2
            ws_summary.cell(row=row, column=1, value=q.get("id", f"Q{qi+1}"))
            ws_summary.cell(row=row, column=2, value=q.get("question", ""))
            ws_summary.cell(row=row, column=3, value=str(q.get("正确回答", "")))
            ws_summary.cell(row=row, column=4, value=str(q.get("point", "")))
            ws_summary.cell(row=row, column=5, value=str(q.get("必答点", "")))
            ws_summary.cell(row=row, column=6, value=str(q.get("选答点", "")))
            ws_summary.cell(row=row, column=7, value=str(q.get("参考文件", "")))

        chain = build_rag_chain(config.score_model)

        for pi, pipe in enumerate(pipelines):
            print(f"[{pi+1}/{len(pipelines)}] {pipe['label']}")
            self._run_one_pipeline(pipe, test_set, ws_summary, pi, wb, chain)

        result_path = out_dir / "result.xlsx"
        wb.save(result_path)
        print(f"\n结果保存至: {result_path}")

        # 评分
        from pipeline.scoring import score_eval_output
        scored_path = out_dir / "scored.xlsx"
        strategy_labels = [p["label"] for p in pipelines]
        score_eval_output(str(result_path), str(scored_path), strategy_labels,
                          strict=(config.score_mode == "strict"))

        # 报告
        from pipeline.report import generate_report
        report_path = out_dir / "report.html"
        generate_report(str(scored_path), str(report_path), title=config.name)

        elapsed = time.time() - t_start
        print(f"\n实验完成. 用时: {elapsed/60:.1f} 分钟")
        print(f"输出目录: {out_dir}")

    def _run_one_pipeline(self, pipe: dict, test_set: list[dict],
                          ws_summary, pipe_index: int, wb, chain):
        sname = pipe["label"]
        col = 8 + pipe_index * 3
        top_k = pipe["top_k"]
        kb = self.config.kb_name

        ws_detail = wb.create_sheet(sname[:31])
        ws_detail.cell(row=1, column=1, value="问题ID")
        ws_detail.cell(row=1, column=2, value="问题")
        ws_detail.cell(row=1, column=3, value="正确回答")
        ws_detail.cell(row=1, column=4, value="回答")
        for ki in range(top_k):
            ws_detail.cell(row=1, column=5 + ki * 2, value=f"检索内容{ki+1}")
            ws_detail.cell(row=1, column=6 + ki * 2, value=f"来源{ki+1}")
        ws_detail.cell(row=1, column=5 + top_k * 2, value="召回")
        ws_detail.cell(row=1, column=6 + top_k * 2, value="召回位置")

        def _process_one(qi_q):
            qi, q = qi_q
            question = q["question"]
            docs = strategy_search(
                query=question, kb_name=kb, chunker_name=pipe["chunker"],
                top_k=top_k, retriever=pipe["retriever"],
            )
            if docs:
                answer = chain.invoke({"question": question, "retrieved_docs": docs})
            else:
                answer = "（无检索结果）"
            return qi, str(answer), docs

        results = {}
        total = len(test_set)
        with ThreadPoolExecutor(max_workers=GEN_WORKERS) as executor:
            futures = {executor.submit(_process_one, (qi, q)): qi for qi, q in enumerate(test_set)}
            for future in as_completed(futures):
                qi, answer, docs = future.result()
                results[qi] = (answer, docs)
                print(f"  Q{qi+1}: {len(answer)} chars  [{len(results)}/{total}]")

        for qi, q in enumerate(test_set):
            row = qi + 2
            answer, docs = results[qi]
            correct = str(q.get("正确回答", ""))
            recall, recall_pos = calc_recall(correct, docs)

            ws_summary.cell(row=row, column=col, value=str(answer)[:30000])
            ws_summary.cell(row=row, column=col + 1, value=recall)
            ws_summary.cell(row=row, column=col + 2, value=f"top{recall_pos}" if recall else 0)

            ws_detail.cell(row=row, column=1, value=q.get("id", f"Q{qi+1}"))
            ws_detail.cell(row=row, column=2, value=q["question"])
            ws_detail.cell(row=row, column=3, value=correct[:30000])
            ws_detail.cell(row=row, column=4, value=str(answer)[:30000])
            for di, d in enumerate(docs):
                ws_detail.cell(row=row, column=5 + di * 2, value=str(d.get("content_text", ""))[:30000])
                ws_detail.cell(row=row, column=6 + di * 2, value=str(d.get("source", "")))
            ws_detail.cell(row=row, column=5 + top_k * 2, value=recall)
            ws_detail.cell(row=row, column=6 + top_k * 2, value=f"top{recall_pos}" if recall else 0)

    def _ensure_built(self, pipelines: list[dict]):
        client = get_client()
        collections = [c.name for c in client.get_collections().collections]
        prefix = f"{self.config.kb_name}__"
        existing = {name for name in collections if name.startswith(prefix)}
        needed = set()
        for p in pipelines:
            needed.add(f"{self.config.kb_name}__{p['chunker']}")
        missing = needed - existing
        client.close()

        if missing:
            print(f"以下集合缺失，自动构建: {missing}")
            from builders.orchestrator import build_all
            build_all(self.config.kb_name)

    def _load_test_set(self, path: str) -> list[dict]:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        test_set = []
        for row in range(2, ws.max_row + 1):
            question = str(ws.cell(row=row, column=1).value or "").strip()
            if not question:
                continue
            q = {"id": f"Q{row-1}", "question": question}
            for ci in range(2, ws.max_column + 1):
                hdr = str(ws.cell(row=1, column=ci).value or "").strip()
                val = ws.cell(row=row, column=ci).value
                if hdr and val:
                    q[hdr] = str(val)
            _ALIAS = {"参考答案": "正确回答", "文件来源": "参考文件"}
            for src, dst in _ALIAS.items():
                if src in q and dst not in q:
                    q[dst] = q[src]
            test_set.append(q)
        return test_set


# ── 纯召回率实验 ──

@dataclass
class RecallConfig:
    name: str = ""
    kb_name: str = DEFAULT_KB
    test_set_path: str = ""
    output_dir: str = ""

    variable: str = "chunker"
    variable_values: list[str] = field(default_factory=list)

    chunker: str = "fixed_1100"
    retriever: str = "hybrid"
    reranker: str = "direct"
    top_k: int = DEFAULT_TOP_K
    top_k_values: list[int] = field(default_factory=lambda: [5, 10])

    def __post_init__(self):
        if not self.name:
            var_label = {"chunker": "多策略", "retriever": "检索器", "reranker": "重排器"}
            self.name = f"{var_label.get(self.variable, self.variable)}召回对比_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        if not self.output_dir:
            self.output_dir = f"output/{self.kb_name}/{self.name}"


class RecallRunner:
    """纯召回率实验：检索 → 召回计算 → 输出，无 LLM 调用。"""

    def __init__(self, config: RecallConfig):
        self.config = config

    def run(self):
        config = self.config
        out_dir = Path(config.output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        (out_dir / "config.json").write_text(
            json.dumps(asdict(config), ensure_ascii=False, indent=2), encoding="utf-8")

        test_set = self._load_test_set(config.test_set_path)
        pipelines = expand_recall_pipelines(config)

        print(f"召回实验: {config.name}")
        print(f"知识库: {config.kb_name}")
        print(f"对比维度: {config.variable}")
        print(f"组合数: {len(pipelines)}  |  测试题数: {len(test_set)}")
        print(f"Top-K: {config.top_k_values}")
        print()

        self._ensure_built(pipelines)

        t_start = time.time()

        for top_k in config.top_k_values:
            self._run_one_topk(top_k, test_set, pipelines, out_dir)

        elapsed = time.time() - t_start
        print(f"\n召回实验完成. 总用时: {elapsed:.1f} 秒")
        print(f"输出目录: {out_dir}")

    def _run_one_topk(self, top_k: int, test_set: list[dict],
                      pipelines: list[dict], out_dir: Path):
        """运行单个 top-k 实验，写入 {out_dir}/top{top_k}/。"""
        config = self.config
        tk_dir = out_dir / f"top{top_k}"
        tk_dir.mkdir(parents=True, exist_ok=True)

        print(f"\n{'─' * 50}")
        print(f"Top-{top_k}")
        print(f"{'─' * 50}")

        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "汇总"
        ws_summary.cell(row=1, column=1, value="问题ID")
        ws_summary.cell(row=1, column=2, value="问题")
        ws_summary.cell(row=1, column=3, value="正确回答")
        ws_summary.cell(row=1, column=4, value="参考文件")
        col = 5
        for pipe in pipelines:
            ws_summary.cell(row=1, column=col, value=f"{pipe['label']}_召回")
            ws_summary.cell(row=1, column=col + 1, value=f"{pipe['label']}_召回位置")
            ws_summary.cell(row=1, column=col + 2, value=f"{pipe['label']}_合并召回")
            ws_summary.cell(row=1, column=col + 3, value=f"{pipe['label']}_拼接召回")
            ws_summary.cell(row=1, column=col + 4, value=f"{pipe['label']}_文件召回")
            col += 5

        for qi, q in enumerate(test_set):
            row = qi + 2
            ws_summary.cell(row=row, column=1, value=q.get("id", f"Q{qi+1}"))
            ws_summary.cell(row=row, column=2, value=q.get("question", ""))
            ws_summary.cell(row=row, column=3, value=str(q.get("正确回答", "")))
            ws_summary.cell(row=row, column=4, value=str(q.get("参考文件", "")))

        recall_stats = {}
        for pi, pipe in enumerate(pipelines):
            sname = pipe["label"]
            print(f"[{pi+1}/{len(pipelines)}] {sname}")

            ws_detail = wb.create_sheet(sname[:31])
            ws_detail.cell(row=1, column=1, value="问题ID")
            ws_detail.cell(row=1, column=2, value="问题")
            ws_detail.cell(row=1, column=3, value="正确回答")
            for ki in range(top_k):
                ws_detail.cell(row=1, column=4 + ki * 3, value=f"检索内容{ki+1}")
                ws_detail.cell(row=1, column=5 + ki * 3, value=f"来源{ki+1}")
                ws_detail.cell(row=1, column=6 + ki * 3, value=f"分数{ki+1}")
            ws_detail.cell(row=1, column=4 + top_k * 3, value="召回")
            ws_detail.cell(row=1, column=5 + top_k * 3, value="召回位置")
            ws_detail.cell(row=1, column=6 + top_k * 3, value="合并召回")
            ws_detail.cell(row=1, column=7 + top_k * 3, value="拼接召回")

            hits = 0
            merged_hits = 0
            concat_hits = 0
            file_hits = 0
            hit_positions = []
            total_questions = 0

            # 并发检索：每个问题独立调用 embedding API + Qdrant
            with ThreadPoolExecutor(max_workers=SEARCH_WORKERS) as pool:
                futures = {}
                for qi, q in enumerate(test_set):
                    f = pool.submit(
                        strategy_search,
                        query=q["question"],
                        kb_name=config.kb_name,
                        chunker_name=pipe["chunker"],
                        top_k=top_k,
                        retriever=pipe["retriever"],
                    )
                    futures[f] = qi

                # 预分配结果槽位，按原始顺序写入
                results = [None] * len(test_set)
                for f in as_completed(futures):
                    qi = futures[f]
                    try:
                        docs = f.result()
                    except Exception as e:
                        print(f"  [Q{qi+1}] 检索失败: {e}")
                        docs = []
                    results[qi] = docs

            for qi, q in enumerate(test_set):
                row = qi + 2
                question = q["question"]
                correct = str(q.get("正确回答", ""))
                ref_file = str(q.get("参考文件", ""))
                docs = results[qi] or []
                recall, recall_pos = calc_recall(correct, docs)
                recall_m, _ = calc_recall_merged(correct, docs)
                recall_c, _ = calc_recall_concat(correct, docs)
                file_recall, file_pos = calc_file_recall(ref_file, docs, top_k)

                total_questions += 1
                if recall:
                    hits += 1
                    hit_positions.append(recall_pos)
                if recall_m:
                    merged_hits += 1
                if recall_c:
                    concat_hits += 1
                if file_recall:
                    file_hits += 1

                col = 5 + pi * 5
                ws_summary.cell(row=row, column=col, value=recall)
                ws_summary.cell(row=row, column=col + 1, value=f"top{recall_pos}" if recall else 0)
                ws_summary.cell(row=row, column=col + 2, value=recall_m)
                ws_summary.cell(row=row, column=col + 3, value=recall_c)
                ws_summary.cell(row=row, column=col + 4, value=file_recall)

                ws_detail.cell(row=row, column=1, value=q.get("id", f"Q{qi+1}"))
                ws_detail.cell(row=row, column=2, value=question)
                ws_detail.cell(row=row, column=3, value=correct[:30000])
                for di, d in enumerate(docs):
                    ws_detail.cell(row=row, column=4 + di * 3, value=str(d.get("content_text", ""))[:30000])
                    ws_detail.cell(row=row, column=5 + di * 3, value=str(d.get("source", "")))
                    ws_detail.cell(row=row, column=6 + di * 3, value=round(float(d.get("score") or 0), 5))
                ws_detail.cell(row=row, column=4 + top_k * 3, value=recall)
                ws_detail.cell(row=row, column=5 + top_k * 3, value=f"top{recall_pos}" if recall else 0)
                ws_detail.cell(row=row, column=6 + top_k * 3, value=recall_m)
                ws_detail.cell(row=row, column=7 + top_k * 3, value=recall_c)

            recall_rate = round(hits / total_questions * 100, 1) if total_questions else 0
            merged_rate = round(merged_hits / total_questions * 100, 1) if total_questions else 0
            concat_rate = round(concat_hits / total_questions * 100, 1) if total_questions else 0
            file_rate = round(file_hits / total_questions * 100, 1) if total_questions else 0
            avg_pos = round(sum(hit_positions) / len(hit_positions), 2) if hit_positions else 0
            recall_stats[sname] = (recall_rate, merged_rate, concat_rate, avg_pos, file_rate)
            print(f"  召回率: {recall_rate}% ({hits}/{total_questions}), 合并召回率: {merged_rate}% ({merged_hits}/{total_questions}), 拼接召回率: {concat_rate}% ({concat_hits}/{total_questions}), 文件召回率: {file_rate}% ({file_hits}/{total_questions}), 平均命中位置: top{avg_pos}")

        # 综合召回 sheet
        ws_score = wb.create_sheet("综合召回")
        ws_score.cell(row=1, column=1, value="策略").font = openpyxl.styles.Font(bold=True)
        ws_score.cell(row=1, column=2, value="召回率(%)").font = openpyxl.styles.Font(bold=True)
        ws_score.cell(row=1, column=3, value="合并召回率(%)").font = openpyxl.styles.Font(bold=True)
        ws_score.cell(row=1, column=4, value="拼接召回率(%)").font = openpyxl.styles.Font(bold=True)
        ws_score.cell(row=1, column=5, value="平均命中位置").font = openpyxl.styles.Font(bold=True)
        ws_score.cell(row=1, column=6, value="文件召回率(%)").font = openpyxl.styles.Font(bold=True)
        for ri, (sname, (rate, merged_rate, concat_rate, pos, file_rate)) in enumerate(recall_stats.items(), start=2):
            ws_score.cell(row=ri, column=1, value=sname)
            ws_score.cell(row=ri, column=2, value=rate)
            ws_score.cell(row=ri, column=3, value=merged_rate)
            ws_score.cell(row=ri, column=4, value=concat_rate)
            ws_score.cell(row=ri, column=5, value=f"top{pos}" if pos else "--")
            ws_score.cell(row=ri, column=6, value=file_rate)

        result_path = tk_dir / "recall.xlsx"
        wb.save(result_path)
        print(f"\nTop-{top_k} 结果保存至: {result_path}")

        from pipeline.report import generate_recall_report
        report_path = tk_dir / "report.html"
        generate_recall_report(str(result_path), str(report_path),
                              title=f"{config.name} (Top-{top_k})")

    def _ensure_built(self, pipelines: list[dict]):
        client = get_client()
        collections = [c.name for c in client.get_collections().collections]
        prefix = f"{self.config.kb_name}__"
        existing = {name for name in collections if name.startswith(prefix)}
        needed = set()
        for p in pipelines:
            needed.add(f"{self.config.kb_name}__{p['chunker']}")
        missing = needed - existing
        client.close()

        if missing:
            print(f"以下集合缺失，自动构建: {missing}")
            from builders.orchestrator import build_all
            build_all(self.config.kb_name)

    def _load_test_set(self, path: str) -> list[dict]:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        test_set = []
        for row in range(2, ws.max_row + 1):
            question = str(ws.cell(row=row, column=1).value or "").strip()
            if not question:
                continue
            q = {"id": f"Q{row-1}", "question": question}
            for ci in range(2, ws.max_column + 1):
                hdr = str(ws.cell(row=1, column=ci).value or "").strip()
                val = ws.cell(row=row, column=ci).value
                if hdr and val:
                    q[hdr] = str(val)
            _ALIAS = {"参考答案": "正确回答", "文件来源": "参考文件"}
            for src, dst in _ALIAS.items():
                if src in q and dst not in q:
                    q[dst] = q[src]
            test_set.append(q)
        return test_set


def expand_recall_pipelines(config: RecallConfig) -> list[dict]:
    pipelines = []
    for value in config.variable_values:
        pipe = {
            "label": value,
            "chunker": value if config.variable == "chunker" else config.chunker,
            "retriever": value if config.variable == "retriever" else config.retriever,
            "reranker": value if config.variable == "reranker" else config.reranker,
            "top_k": config.top_k,
        }
        pipelines.append(pipe)
    return pipelines